"""Lectura de fuentes locales. No descarga datos ni altera los originales."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import struct
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from modelo import INICIO, FIN, coordenadas_validas, numero, resumir_serie


def normalizar(texto):
    return " ".join("".join(c for c in unicodedata.normalize("NFD", str(texto))
                           if unicodedata.category(c) != "Mn").lower().split())


class Fuentes:
    def __init__(self, raiz):
        self.raiz = Path(raiz).resolve()
        self.archivos = {}
        self.avisos = []
        self.catalogos = {}

    def registrar(self, ruta):
        ruta = Path(ruta)
        if ruta not in self.archivos:
            try:
                relativa = ruta.relative_to(self.raiz).as_posix()
            except ValueError:
                if ruta.parent != Path(__file__).resolve().parent:
                    raise ValueError("Archivo fuera del repositorio: "+str(ruta))
                relativa = "Pruebas/"+ruta.name
            self.archivos[ruta] = {
                "ruta": relativa,
                "bytes": ruta.stat().st_size,
                "sha256": hashlib.sha256(ruta.read_bytes()).hexdigest(),
            }

    def json(self, ruta):
        ruta = Path(ruta)
        self.registrar(ruta)
        return json.loads(ruta.read_text(encoding="utf-8-sig"))

    def carpeta_dmc(self, nombre):
        base = self.raiz / "FUENTES" / "CONFORT TERMICO"
        ruta = base / nombre
        if ruta.is_dir():
            return ruta
        # Windows elimina puntos finales en carpetas al extraer ZIP de GitHub.
        alternativas = [p for p in base.iterdir() if p.is_dir()
                        and p.name.rstrip(". ") == nombre.rstrip(". ")]
        return alternativas[0] if len(alternativas) == 1 else ruta

    def clima(self, config):
        import openpyxl
        estaciones = []
        for meta in config["estaciones"]:
            estacion = {**meta, "resumenes": {}}
            carpeta = self.carpeta_dmc(meta["carpeta"])
            if not carpeta.is_dir():
                self.avisos.append("Carpeta de estación ausente: "+meta["carpeta"])
            observaciones = defaultdict(list)
            errores = Counter()
            archivos_por_componente = defaultdict(list)
            for archivo in sorted(carpeta.glob("*.xlsx")):
                self.registrar(archivo)
                wb = openpyxl.load_workbook(archivo, read_only=True, data_only=False)
                try:
                    hojas = []
                    for hoja in wb.worksheets:
                        primera = next(hoja.iter_rows(), None)
                        if primera and len(primera) == 3:
                            rotulos = [normalizar(c.value) for c in primera]
                            if rotulos[:2] == ["fecha", "hora (utc)"] and rotulos[2] in ("ts (°c)", "hr (%)"):
                                hojas.append(hoja)
                    if len(hojas) != 1:
                        raise ValueError("Se esperaba una hoja inequívoca de datos meteorológicos")
                    filas = hojas[0].iter_rows()
                    cabecera = next(filas, None)
                    if not cabecera or len(cabecera) != 3:
                        raise ValueError("Se esperaban Fecha, Hora (UTC) y valor")
                    etiquetas = [normalizar(c.value) for c in cabecera]
                    if etiquetas[:2] != ["fecha", "hora (utc)"]:
                        raise ValueError("Cabecera de fecha/hora no reconocida")
                    tercera = etiquetas[2]
                    if tercera == "ts (°c)":
                        componente = "temperatura"
                    elif tercera == "hr (%)":
                        componente = "humedad"
                    else:
                        raise ValueError("Variable o unidad no reconocida: "+str(cabecera[2].value))
                    archivos_por_componente[componente].append(
                        archivo.relative_to(self.raiz).as_posix()+"#"+hojas[0].title)
                    for fila in filas:
                        if all(c.value is None for c in fila):
                            continue
                        if any(c.data_type == "f" for c in fila):
                            errores[componente] += 1
                            continue
                        try:
                            dia, hora, v = (c.value for c in fila)
                            if isinstance(dia, datetime):
                                dia = dia.date()
                            elif not isinstance(dia, date):
                                dia = datetime.strptime(str(dia), "%d-%m-%Y").date()
                            if isinstance(hora, datetime):
                                hora = hora.time()
                            elif not isinstance(hora, time):
                                hora = time.fromisoformat(str(hora))
                            observaciones[componente].append((datetime.combine(dia, hora), v))
                        except (ValueError, TypeError):
                            errores[componente] += 1
                except (ValueError, StopIteration) as e:
                    self.avisos.append(f"{archivo.relative_to(self.raiz)}: {e}")
                finally:
                    wb.close()
            for componente in ("temperatura", "humedad"):
                resumen = resumir_serie(observaciones[componente], componente)
                resumen["errores_lectura"] = errores[componente]
                resumen["archivos"] = archivos_por_componente[componente]
                if componente in meta.get("componentes_en_revision", []):
                    resumen.update(puntaje_100=None, valor_medio=None,
                                   estado="En revisión de variable/unidad/validez")
                estacion["resumenes"][componente] = resumen
            estaciones.append(estacion)
        self.catalogos["dmc"] = {"estaciones": len(estaciones),
                                 "archivos_excel": sum(len(s["resumenes"][c]["archivos"])
                                     for s in estaciones for c in s["resumenes"])}
        return estaciones

    def aire(self, config):
        base = self.raiz / "FUENTES" / "fuentes climaticas usadas"
        estaciones = []
        for meta in config["estaciones"]:
            estacion = {**meta, "resumenes": {}}
            for serie in meta["series"]:
                componente = serie["componente"]
                archivo = base / serie["archivo"]
                self.registrar(archivo)
                datos, conteos, errores = leer_csv_aire(archivo)
                confirmado = (
                    serie.get("metadatos_confirmados") is True
                    and serie.get("unidad") == "ug/m3"
                    and serie.get("agregacion") in ("horaria", "diaria")
                    and bool(serie.get("huso_horario"))
                    and bool(serie.get("evidencia"))
                )
                if confirmado:
                    resumen = resumir_serie(datos, componente, serie["agregacion"])
                else:
                    resumen = {"puntaje_100": None, "valor_medio": None,
                               "estado": "Metadatos de exportación pendientes",
                               "meses_usados": [], "meses": []}
                resumen.update(
                    archivo=archivo.relative_to(self.raiz).as_posix(),
                    conteos_validacion_en_periodo=conteos, errores_csv=errores,
                    metadatos=serie,
                )
                estacion["resumenes"][componente] = resumen
            for componente in ("mp25", "o3"):
                estacion["resumenes"].setdefault(componente, {
                    "puntaje_100": None, "valor_medio": None,
                    "estado": "Sin archivo", "meses_usados": [], "meses": []})
            estaciones.append(estacion)
        self.catalogos["aire"] = {
            "estaciones": len(estaciones),
            "csv": sum(len(s["series"]) for s in config["estaciones"]),
            "nota": "Sin metadatos confirmados no se calcula aire; no validados excluidos.",
        }
        return estaciones

    def salud(self):
        archivo = (self.raiz / "FUENTES" / "fuentes cercania a establecimiento de salud"
                   / "ESTABLECIMIENTOS_12062026_DES.zip")
        self.registrar(archivo)
        with zipfile.ZipFile(archivo) as z:
            prj = next(n for n in z.namelist() if n.lower().endswith(".prj"))
            proyeccion = z.read(prj).decode("utf-8-sig")
            if "WGS_1984" not in proyeccion and "WGS 84" not in proyeccion:
                raise ValueError("El catálogo de salud no acredita WGS84")
            nombre = next(n for n in z.namelist() if n.lower().endswith(".dbf"))
            filas = leer_dbf(z.read(nombre))
        tipos = {
            "Hospital", "Clínica", "Centro de Salud Privado",
            "Centro de Salud Familiar (CESFAM)",
            "Centro Comunitario de Salud Familiar (CECOSF)",
            "Posta de Salud Rural (PSR)", "Consultorio General Urbano (CGU)",
            "Consultorio General Rural (CGR)",
            "Servicio de Atención Primaria de Urgencia (SAPU)",
            "Servicio de Atención Primaria de Urgencia de Alta Resolutividad (SAR)",
            "Servicio de Urgencia Rural (SUR)",
        }
        tipos = {normalizar(x) for x in tipos}
        por_id = defaultdict(list)
        filtrados = 0
        sin_coord = 0
        for r in filas:
            if (normalizar(r["TIPO"]) not in tipos
                    or normalizar(r["ESTADO"]) != "vigente en operacion habitual"
                    or normalizar(r["PRESTADOR"]) not in ("publico", "privado")):
                continue
            filtrados += 1
            lat, lon = numero(r["LATITUD"]), numero(r["LONGITUD"])
            if not coordenadas_validas(lat, lon):
                sin_coord += 1
                continue
            try:
                codigo = str(int(Decimal(r["COD_VIG"])))
            except (ValueError, InvalidOperation):
                continue
            por_id[codigo].append({
                "codigo": codigo, "nombre": r["NOMBRE"], "tipo": r["TIPO"],
                "prestador": r["PRESTADOR"], "latitud": lat, "longitud": lon,
            })
        catalogo = []
        conflictos = 0
        for registros in por_id.values():
            if any(r != registros[0] for r in registros[1:]):
                conflictos += 1
            else:
                catalogo.append(registros[0])
        self.catalogos["salud"] = {
            "archivo": archivo.relative_to(self.raiz).as_posix(), "registros": len(filas),
            "filtrados_por_tipo_estado_prestador": filtrados,
            "coordenadas_invalidas": sin_coord, "ids_en_conflicto": conflictos,
            "establecimientos_admisibles": len(catalogo), "edicion": "junio 2026",
        }
        return catalogo

    def paraderos(self):
        base = self.raiz / "HABITA_GTFS_Movilidad_v1.2.1_GUI" / "data"
        manifest = self.json(base / "manifest.json")
        partes = []
        for nombre in manifest["parts"]:
            p = base / nombre
            self.registrar(p)
            partes.append(p.read_bytes())
        contenido = b"".join(partes)
        if len(contenido) != manifest["size"]:
            raise ValueError("Tamaño del GTFS reconstruido incompatible con manifest.json")
        if hashlib.sha256(contenido).hexdigest() != manifest["sha256"]:
            raise ValueError("SHA256 del GTFS incompatible con manifest.json")
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            nombres = {Path(n).name: n for n in z.namelist()}
            def leer(nombre):
                return csv.DictReader(io.StringIO(
                    z.read(nombres[nombre]).decode("utf-8-sig")))
            feed = next(leer("feed_info.txt"))
            rutas = {r["route_id"] for r in leer("routes.txt") if r["route_type"] == "3"}
            viajes = {r["trip_id"] for r in leer("trips.txt") if r["route_id"] in rutas}
            abordables = {r["stop_id"] for r in leer("stop_times.txt")
                          if r["trip_id"] in viajes and r.get("pickup_type", "") != "1"}
            grupos = defaultdict(list)
            for r in leer("stops.txt"):
                lat, lon = numero(r["stop_lat"]), numero(r["stop_lon"])
                if (r.get("location_type", "") not in ("", "0")
                        or r["stop_id"] not in abordables
                        or not coordenadas_validas(lat, lon)):
                    continue
                grupos[r["stop_id"]].append({
                    "codigo": r["stop_id"], "nombre": r["stop_name"],
                    "latitud": lat, "longitud": lon,
                })
            catalogo = [rs[0] for rs in grupos.values()
                        if all(r == rs[0] for r in rs)]
        self.catalogos["paraderos"] = {
            "feed": feed, "sha256": manifest["sha256"],
            "paraderos_admisibles": len(catalogo),
            "filtro": "location_type vacío/0; bus route_type 3; presencia abordable en stop_times",
            "alcance": "Ubicación física; no se evalúan frecuencia ni viajes activos.",
        }
        return catalogo


def leer_dbf(b):
    """DBF del catálogo MINSAL; conserva textos, interpreta campos en el filtro."""
    n = struct.unpack_from("<I", b, 4)[0]
    cabecera, largo = struct.unpack_from("<HH", b, 8)
    campos = []
    pos = 32
    while pos < cabecera and b[pos] != 13:
        d = b[pos:pos+32]
        campos.append((d[:11].split(b"\0")[0].decode("ascii"), d[16]))
        pos += 32
    filas = []
    for i in range(n):
        registro = b[cabecera+i*largo:cabecera+(i+1)*largo]
        if len(registro) != largo:
            raise ValueError("DBF truncado")
        if registro[:1] == b"*":
            continue
        r, off = {}, 1
        for nombre, ancho in campos:
            raw = registro[off:off+ancho].strip()
            off += ancho
            try:
                r[nombre] = raw.decode("utf-8")
            except UnicodeDecodeError:
                r[nombre] = raw.decode("cp1252")
        filas.append(r)
    return filas


def leer_csv_aire(archivo):
    conteos = Counter()
    candidatos = defaultdict(list)
    errores = 0
    texto = Path(archivo).read_bytes().decode("utf-8-sig")
    filas = csv.reader(io.StringIO(texto), delimiter=";")
    cab = next(filas)
    if cab[:5] != ["FECHA (YYMMDD)", "HORA (HHMM)", "Registros validados", "Registros preliminares", "Registros no validados"]:
        raise ValueError("Cabecera SINCA no reconocida: "+str(archivo))
    for r in filas:
        if not any(r):
            continue
        try:
            dia = datetime.strptime(r[0], "%y%m%d").date()
            hora = datetime.strptime(r[1].zfill(4), "%H%M").time()
            marca = datetime.combine(dia, hora)
        except (ValueError, IndexError):
            errores += 1
            continue
        if not INICIO <= dia <= FIN:
            continue
        for i, etiqueta in ((2, "validados"), (3, "preliminares"), (4, "no_validados")):
            if len(r) > i and r[i].strip():
                conteos[etiqueta] += 1
        # La columna de mejor estado disponible tiene prioridad incluso si es cero.
        for prioridad in (2, 3):
            if len(r) > prioridad and r[prioridad].strip():
                candidatos[marca].append((prioridad, numero(r[prioridad])))
                break
    datos = []
    for marca, valores in candidatos.items():
        prioridad = min(x[0] for x in valores)
        # Resumir_serie elimina conflictos de mismo estado/misma marca.
        elegidos = [x for x in valores if x[0] == prioridad]
        conteos["seleccionados_"+("validados" if prioridad == 2 else "preliminares")] += len(elegidos)
        datos.extend((marca, x[1]) for x in elegidos)
    return datos, dict(conteos), errores
